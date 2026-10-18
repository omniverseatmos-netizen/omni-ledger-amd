from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from PIL import Image, ImageEnhance
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'PITCH_DECK.pptx'
ASSET = ROOT / 'assets' / 'omni-ledger-cover.png'
BG_IMG = ROOT / 'assets' / 'omni-ledger-cover-16x9.jpg'

# Build a clean 16:9 crop for the deck while preserving the generated artwork.
im = Image.open(ASSET).convert('RGB')
w, h = im.size
target_ratio = 16/9
if w/h < target_ratio:
    new_h = int(w/target_ratio)
    top = (h-new_h)//2
    im = im.crop((0, top, w, top+new_h))
else:
    new_w = int(h*target_ratio)
    left = (w-new_w)//2
    im = im.crop((left, 0, left+new_w, h))
im = im.resize((1920,1080), Image.Resampling.LANCZOS)
im.save(BG_IMG, quality=90, optimize=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

C = {
    'bg': '080D16', 'panel': '0D1623', 'panel2': '101D2D', 'panel3':'142338',
    'line':'203247', 'white':'F2F6FB', 'muted':'91A1B7', 'muted2':'6F839B',
    'teal':'6FE6C5', 'blue':'80B7FF', 'amber':'FFCA80', 'rose':'FF8093', 'green':'7DE2BD',
    'dark':'071A1A', 'pale':'DCE7F3'
}

def rgb(hexv): return RGBColor.from_string(hexv)

def bg(slide, color=C['bg']):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = rgb(color)

def rect(slide,x,y,w,h,fill=None,line=None,radius=True,trans=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill or C['panel']); shape.fill.transparency = trans
    shape.line.fill.solid(); shape.line.color.rgb = rgb(line or C['line']); shape.line.width = Pt(0.7)
    if hasattr(shape, 'adjustments') and radius:
        try: shape.adjustments[0] = 0.12
        except: pass
    return shape

def txt(slide, text, x,y,w,h, size=14, color=C['white'], bold=False, font='Arial', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02, char_spacing=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin); tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, line in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(2)
        for r in p.runs:
            r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
            if char_spacing is not None:
                try: r._r.get_or_add_rPr().set('spc', str(char_spacing))
                except: pass
    return shape

def label(slide, text, x, y, w=4, color=C['teal']):
    return txt(slide, text.upper(), x,y,w,.22, size=8, color=color, bold=True, char_spacing=120)

def title(slide, eyebrow, heading, sub=None):
    label(slide, eyebrow, .65, .42, 8)
    txt(slide, heading, .65, .76, 12.0, .62, size=26, bold=True)
    if sub: txt(slide, sub, .68, 1.43, 11.7, .48, size=10, color=C['muted'])

def footer(slide, idx):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(7.12), Inches(12.03), Inches(.008))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(C['line']); line.line.fill.background()
    txt(slide, 'OmNi-Ledger  /  Technical evaluation prototype', .68, 7.18, 8, .18, size=7, color=C['muted2'])
    txt(slide, f'{idx:02d}  /  08', 11.75, 7.17, .88, .2, size=7, color=C['muted2'], align=PP_ALIGN.RIGHT)

def card(slide, x,y,w,h, head, body, accent=C['teal'], small=None):
    rect(slide,x,y,w,h,C['panel'],C['line'])
    rect(slide,x,y,.045,h,accent,accent,radius=False)
    txt(slide, head, x+.2,y+.2,w-.4,.35,size=13,bold=True)
    txt(slide, body, x+.2,y+.64,w-.4,h-.88,size=10,color=C['muted'])
    if small: txt(slide, small, x+.2,y+h-.3,w-.4,.16,size=7,color=accent,bold=True)

# Slide 1 — Cover
s = prs.slides.add_slide(blank)
s.shapes.add_picture(str(BG_IMG), 0, 0, width=prs.slide_width, height=prs.slide_height)
# Subtle dark navy panel on left to guarantee readability.
rect(s,0,0,7.25,7.5,C['bg'],C['bg'],radius=False,trans=16)
label(s,'AMD DEVELOPER HACKATHON · ACT III',.78,.72,5.8)
txt(s,'OmNi-Ledger',.78,1.48,6.15,.8,size=38,bold=True)
txt(s,'Deterministic supply-chain control\nfor AI-assisted pharmaceutical intake',.8,2.43,6.0,1.18,size=22,bold=True,color=C['pale'])
txt(s,'AI may extract and propose.\nDeterministic rules decide what is allowed.',.82,4.0,5.45,.76,size=15,color=C['teal'],bold=True)
rect(s,.8,5.2,4.1,.46,C['panel2'],C['line'])
txt(s,'PROPOSED TRACK  ·  REINVENT COMMERCE',.8,5.2,4.1,.46,size=8,color=C['blue'],bold=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
txt(s,'Solo-founder prototype  ·  Egypt  ·  Technical evaluation build',.82,6.7,6.0,.23,size=8,color='BDCCDC')
# Brand mark block
rect(s,11.95,.62,.6,.6,C['panel2'],C['teal'])
txt(s,'OL',12.04,.78,.42,.2,size=13,color=C['teal'],bold=True,align=PP_ALIGN.CENTER)

# Slide 2 — Problem
s = prs.slides.add_slide(blank); bg(s); title(s,'01 / Problem','An alert is not a control.','Pharmaceutical supply chains need rules that remain effective after an invoice has been read and stock has entered the warehouse.')
card(s,.65,2.2,3.82,2.12,'01  Wrong item or batch','A plausible document extraction can still refer to the wrong SKU, lot, expiry, or quantity. Human review alone can miss a mismatch. ',C['amber'],'IDENTITY + SCHEMA VALIDATION')
card(s,4.75,2.2,3.82,2.12,'02  Expired stock','If an expired batch becomes available inventory, the error can travel into allocation, picking, packing, and dispatch. ',C['rose'],'DETERMINISTIC HARD BLOCK')
card(s,8.85,2.2,3.82,2.12,'03  Recall after dispatch','A recall must identify remaining stock and affected recorded customer orders—not just change a red status badge. ',C['blue'],'BATCH TRACE + AUDIT')
rect(s,.65,4.6,12.02,1.9,C['panel2'],C['line'])
label(s,'The design question',.9,4.82,3)
txt(s,'How can AI help process invoices without becoming the authority that changes critical inventory?',.9,5.15,10.9,.82,size=19,bold=True)
txt(s,'The answer is a clear trust boundary between probabilistic extraction and deterministic execution.',.9,6.08,10.9,.28,size=10,color=C['muted'])
footer(s,2)

# Slide 3 — Principle
s = prs.slides.add_slide(blank); bg(s); title(s,'02 / Approach','AI reads. Rules decide. Evidence proves.','One connected control path—from document fields to warehouse state transitions and verifiable records.')
steps = [
    ('01','EXTRACT','Gemma on AMD','Untrusted structured fields',C['blue']),
    ('02','VALIDATE','Strict policy gate','Dates · SKU · quantity · authority',C['teal']),
    ('03','CONTROL','Warehouse core','QC · FEFO · scan · recall',C['amber']),
    ('04','PROVE','Audit evidence','SHA-256 chain · state seal · Merkle root',C['green'])
]
xs = [.65,3.77,6.89,10.01]
for i,(n,k,h,b,a) in enumerate(steps):
    rect(s,xs[i],2.42,2.68,2.25,C['panel'],C['line'])
    label(s,n+' / '+k,xs[i]+.2,2.65,2.25,a)
    txt(s,h,xs[i]+.2,3.18,2.25,.58,size=16,bold=True)
    txt(s,b,xs[i]+.2,3.92,2.25,.52,size=9,color=C['muted'])
    if i<3:
        line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xs[i]+2.7), Inches(3.55), Inches(xs[i+1]-.1), Inches(3.55))
        line.line.color.rgb=rgb(C['teal']); line.line.width=Pt(1.2)
rect(s,.65,5.15,12.02,.92,'0D1C29','203A50')
txt(s,'Core invariant',.88,5.15,1.45,.92,size=9,color=C['teal'],bold=True,valign=MSO_ANCHOR.MIDDLE)
txt(s,'The model never writes inventory balances directly.',2.4,5.15,10.1,.92,size=16,bold=True,valign=MSO_ANCHOR.MIDDLE)
txt(s,'Live inference runs in the secure Node CLI. The public judge page runs deterministic local scenarios.',.9,6.32,10.9,.28,size=9,color=C['muted'])
footer(s,3)

# Slide 4 — Controls
s = prs.slides.add_slide(blank); bg(s); title(s,'03 / Implemented controls','Operational rules—not UI promises.','The judge edition exposes controls implemented in the shared domain core. Each visible decision is reproducible in the local scenario.')
card(s,.65,2.16,5.82,1.45,'Expiry & receiving gate','An expired batch is rejected before it can become available stock; posting follows the QC workflow.',C['rose'],'NEGATIVE CONTROL')
card(s,6.82,2.16,5.82,1.45,'FEFO & scan verification','Allocation uses eligible stock; picking verifies both item identity and the allocated batch.',C['blue'],'ORDER STATE MACHINE')
card(s,.65,3.92,5.82,1.45,'Recall & quarantine','A documented recall requires separate initiator and approver identities, then moves available stock in the selected warehouse into quarantine.',C['amber'],'RECALLED_BATCH GUARDS')
card(s,6.82,3.92,5.82,1.45,'Audit integrity','Ledger movements and audit events are hash-linked; state-seal verification can reveal inconsistent data.',C['teal'],'TAMPER-EVIDENT · NOT A LEGAL SIGNATURE')
rect(s,.65,5.78,12.02,.57,C['panel2'],C['line'])
txt(s,'The same core powers the interactive judge demo and the live-CLI post-extraction decision.',.85,5.78,11.6,.57,size=10,color=C['pale'],bold=True,valign=MSO_ANCHOR.MIDDLE)
footer(s,4)

# Slide 5 — Recall scenario
s = prs.slides.add_slide(blank); bg(s); title(s,'04 / Signature moment','Batch recall: trace backward, control forward.','Fictional demonstration scenario — computed by the packaged domain core, not manually authored dashboard numbers.')
# Three stage cards and connecting routes
stages = [(.65,'SUPPLIER','United Pharmaceutical\nSupply Co.','100 boxes received',C['blue']), (4.78,'WAREHOUSE','WH-OCT6 · QUARANTINE','70 boxes held',C['rose']), (8.91,'CUSTOMER','Al-Shifa Pharmacy\n6th of October','30 boxes dispatched',C['teal'])]
for x,head,body,foot,a in stages:
    rect(s,x,2.1,3.55,1.45,C['panel'],C['line'])
    label(s,head,x+.2,2.33,2.8,a)
    txt(s,body,x+.2,2.67,3.1,.45,size=13,bold=True)
    txt(s,foot,x+.2,3.25,3.1,.18,size=8,color=a,bold=True)
for x1,x2 in [(4.23,4.65),(8.36,8.78)]:
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(2.86), Inches(x2), Inches(2.86)); ln.line.color.rgb=rgb(C['teal']); ln.line.width=Pt(1.3)
rect(s,.65,3.92,7.25,1.47,'261823','633044')
label(s,'POST-RECALL PICK ATTEMPT',.9,4.16,4,C['rose'])
txt(s,'HARD BLOCK',.9,4.48,2.65,.52,size=26,bold=True,color=C['rose'])
txt(s,'RECALLED_BATCH',3.55,4.56,3.2,.24,size=11,bold=True,color=C['pale'])
txt(s,'The previously allocated open order cannot be picked.',.9,5.04,6.2,.2,size=9,color='D1AAB5')
rect(s,8.2,3.92,4.27,1.47,C['panel2'],C['line'])
label(s,'RECALL REFERENCE',8.45,4.16,2.8,C['amber'])
txt(s,'DEMO-RECALL-017',8.45,4.52,3.6,.31,size=16,bold=True)
txt(s,'Reason + separated approval recorded',8.45,5.0,3.6,.2,size=8,color=C['muted'])
txt(s,'Scope: batch-level trace through recorded warehouse dispatch. No patient, downstream sub-pharmacy, or individual serial tracking is claimed.',.75,6.07,11.9,.45,size=9,color=C['muted'])
footer(s,5)

# Slide 6 — Evidence
s = prs.slides.add_slide(blank); bg(s); title(s,'05 / Verifiability','Trust is a result you can inspect.','The report connects the business decision to the ledger, actor, time, and integrity chain.')
items = [
    ('LEDGER CHAIN','VALID','Movement records are linked through prior hashes.',C['green']),
    ('AUDIT CHAIN','VALID','Control events preserve their own event sequence.',C['blue']),
    ('STATE SEAL','VERIFIED','State hash and Merkle root are checked against the sample state.',C['teal']),
]
for i,(head,val,desc,a) in enumerate(items):
    x=.65+i*4.1
    rect(s,x,2.25,3.82,1.48,C['panel'],C['line'])
    label(s,head,x+.2,2.49,3,a)
    txt(s,val,x+.2,2.82,3.1,.38,size=23,bold=True,color=a)
    txt(s,desc,x+.2,3.33,3.3,.28,size=8,color=C['muted'])
rect(s,.65,4.15,12.02,1.35,C['panel2'],C['line'])
label(s,'Example evidence fields',.9,4.39,3)
# hash blocks
for i,(k,v) in enumerate([('EVENT HASH','sha256(previousHash + canonicalEvent)'),('STATE HASH','derived from current state'),('MERKLE ROOT','tree root over event/state hashes')]):
    x=.9+i*3.95
    txt(s,k,x,4.8,3.3,.18,size=8,color=C['blue'],bold=True)
    txt(s,v,x,5.08,3.45,.25,size=9,color=C['pale'])
rect(s,.65,5.83,12.02,.58,'221B20','4C3038')
txt(s,'Integrity hashes are not encryption, not a legal signature, and not proof of production security by themselves.',.85,5.83,11.6,.58,size=9,color='FFCFD5',bold=True,valign=MSO_ANCHOR.MIDDLE)
footer(s,6)

# Slide 7 — Boundaries and architecture
s = prs.slides.add_slide(blank); bg(s); title(s,'06 / Engineering honesty','Strong claims have explicit boundaries.','The prototype’s value comes from testable behavior—not from claiming to solve every downstream traceability problem.')
card(s,.65,2.08,5.82,1.55,'Proven in the packaged tests','Valid and expired receipt paths, recall/quarantine, post-recall pick block, recorded customer trace, transfer guard under recall, and ledger/audit verification.',C['green'],'LOCAL AUTOMATED TESTS')
card(s,6.82,2.08,5.82,1.55,'Must be verified with account access','Live Gemma inference through the exact AMD model endpoint, and the current deployed judge URL after publishing this new build.',C['blue'],'LIVE CREDENTIALS / DEPLOYMENT')
card(s,.65,3.98,5.82,1.55,'Outside current scope','Patient-level and downstream sub-pharmacy tracing; individual GS1 serialization; customer returns; multi-user production concurrency.',C['amber'],'NOT IMPLEMENTED')
card(s,6.82,3.98,5.82,1.55,'Not claimed','Regulatory certification, zero hallucinations, unhackable storage, production readiness, or guaranteed competition results.',C['rose'],'NO OVERCLAIMS')
txt(s,'The live request runs in Node/CLI; fixture mode never claims to be live inference.',.75,6.07,11.7,.28,size=10,color=C['teal'],bold=True)
footer(s,7)

# Slide 8 — Mission / close
s = prs.slides.add_slide(blank)
s.shapes.add_picture(str(BG_IMG),0,0,width=prs.slide_width,height=prs.slide_height)
rect(s,0,0,13.333,7.5,C['bg'],C['bg'],radius=False,trans=16)
label(s,'07 / MISSION',.8,.72,3)
txt(s,'AI-assisted operations.\nHuman accountability.\nVerifiable control.',.8,1.35,7.3,2.05,size=31,bold=True)
txt(s,'OmNi-Ledger was developed by a solo founder in Egypt through a phone-first, resource-constrained workflow. The mission is to make critical workflows safer to inspect and harder to bypass—not to hand operational authority to a language model.',.84,3.82,7.0,1.1,size=13,color=C['pale'])
rect(s,.84,5.45,3.3,.5,C['teal'],C['teal'])
txt(s,'LET THE EVIDENCE SPEAK',1.04,5.59,2.85,.2,size=9,color=C['dark'],bold=True,align=PP_ALIGN.CENTER)
txt(s,'OmNi-Ledger  ·  Reinvent Commerce (proposed track)',.85,6.78,7,.2,size=8,color='B7C5D4')
rect(s,10.95,5.65,1.0,1.0,C['panel2'],C['teal'])
txt(s,'OL',11.17,5.95,.58,.28,size=18,color=C['teal'],bold=True,align=PP_ALIGN.CENTER)

prs.core_properties.title = 'OmNi-Ledger — Deterministic Supply-Chain Control'
prs.core_properties.subject = 'AMD Developer Hackathon ACT III · Judge Edition'
prs.core_properties.author = 'OmNi-Ledger project'
prs.core_properties.keywords = 'AMD, Gemma, pharmaceutical supply chain, batch recall, deterministic validation, audit evidence'
prs.core_properties.comments = 'Prototype presentation. See DEMO_SCRIPT.md and SCOPE_AND_LIMITATIONS.md for claims and boundaries.'
prs.save(OUT)
print(f'Created {OUT} ({OUT.stat().st_size} bytes)')
